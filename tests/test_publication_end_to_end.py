"""The repository configuration on a synthetic market with three traps: a bad print on the
S&P 500, a Brent contract roll, and a French 10Y one day late (OAT-Bund spread)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from market_monitor.alerts import load_rules
from market_monitor.export import export_daily_macro, load_layout
from market_monitor.referential import load_referential
from tests.synthetic import synthetic_world

ROOT = Path(__file__).resolve().parents[1] / "config"
AS_OF = "2026-10-01"  # Thursday: Brent generic rolls to the next contract


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    ref = load_referential(ROOT / "instruments.yaml", ROOT / "watchlists.yaml")
    monitor, primary, secondary = synthetic_world(ref, end=AS_OF)
    primary["SPX"] = primary["SPX"].copy()
    primary["SPX"].iloc[-1] *= 1.08                       # bad print on the primary source only
    for data in (primary, secondary):                     # +7 % calendar-spread jump, both sources
        data["BRENT"] = data["BRENT"].copy()
        data["BRENT"].iloc[-1] *= 1.07
        data["OAT_10Y"] = data["OAT_10Y"].iloc[:-1]       # OAT 10Y late by one session
    layout = load_layout(ROOT / "daily_macro.yaml", monitor.referential)
    rules = load_rules(ROOT / "alerts.yaml", monitor.referential)
    out = tmp_path_factory.mktemp("export")
    return export_daily_macro(monitor, layout, out, AS_OF, rules)


def test_bad_print_is_caught_by_the_second_source(result):
    spx = result.data.universe.loc["SPX"]
    assert spx["suspect"] and spx["cross_check"] == "écart"
    assert "FMP" in spx["suspect_reason"]
    assert "SPX" not in result.data.movers.index
    assert result.data.universe.loc["SX5E", "cross_check"] == "ok"


def test_roll_is_flagged_not_reported(result):
    brent = result.data.universe.loc["BRENT"]
    assert brent["roll_1d"] and not brent["suspect"]
    assert "BRENT" not in result.data.movers.index


def test_text_lists_every_point_to_check(result):
    text = result.text.read_text(encoding="utf-8")
    assert "- S&P 500 : " in text and "à vérifier)" in text.split("- S&P 500 : ")[1].split("\n")[0]
    block = text.split("À vérifier avant diffusion")[1]
    assert "S&P 500 : mouvement suspect" in block
    assert "Brent : 1J affecté par un changement de contrat" in block
    assert "Spread OAT-Bund 10 ans : jambes décalées : dernière date commune 30/09" in block


def test_alerts_report_the_data_not_a_market_move(result):
    alerts = result.data.alerts.alerts
    spx = [a for a in alerts if a.instrument_id == "SPX"]
    assert [a.rule_id for a in spx] == ["suspect_data"]
    assert not [a for a in alerts if a.instrument_id == "BRENT" and a.kind.value in ("zscore", "change")]


def test_excel_shades_and_documents_the_lines(result):
    wb = load_workbook(result.excel)
    raw = wb["Données"]
    headers = [c.value for c in raw[1]]
    rows = {r[0].value: r for r in raw.iter_rows(min_row=2)}
    assert rows["SPX"][headers.index("Contrôle 2e source")].value == "écart"
    assert rows["BRENT"][headers.index("Roll 1J")].value == "oui"
    assert "mouvement suspect" in rows["SPX"][headers.index("À vérifier")].value
    macro = wb["Daily macro"]
    spx_cell = next(c for c in macro["A"] if c.value and "$" in str(c.value)
                    and int(str(c.value).rsplit("$", 1)[1]) == rows["SPX"][0].row)
    assert spx_cell.fill.fgColor.rgb.endswith("FCE9C8")
    assert "mouvement suspect" in spx_cell.comment.text


def test_png_is_written(result):
    assert result.png.stat().st_size > 20_000
