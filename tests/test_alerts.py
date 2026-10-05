import json
import math
from datetime import date

import pandas as pd
import pytest
import yaml

from market_monitor.alerts import RuleKind, Severity, evaluate, rules_from_dict
from market_monitor.cli import main
from market_monitor.data.service import MarketDataService
from market_monitor.export import build_daily_macro, render_text, write_excel
from market_monitor.export.layout import layout_from_dict
from market_monitor.monitor import MarketMonitor
from market_monitor.referential import ReferentialError
from tests.test_export import LAYOUT_RAW
from tests.test_history_monitor import REF, _providers

AS_OF = date(2026, 10, 5)


def _row(name, quote="price", unit="pct", asset_class="equity", **values):
    base = {"name": name, "asset_class": asset_class, "quote": quote, "change_unit": unit,
            "level": 100.0, "level_date": pd.Timestamp("2026-10-02"), "stale": False,
            "chg_1d": 0.0, "chg_1w": 0.0, "chg_mtd": 0.0, "chg_ytd": 0.0,
            "z_1d": 0.0, "z_1w": 0.0, "ref_1d": 100.0}
    return {**base, **values}


def _table(**rows):
    return pd.DataFrame.from_dict(rows, orient="index")


def _rules(*entries):
    return rules_from_dict({"rules": list(entries)}, REF)


# ---------------------------------------------------------------------- parsing
def test_parse_rules_and_selectors():
    rules = _rules(
        {"id": "z", "type": "zscore", "instruments": ["watchlist:macro"], "min_abs_z": 2,
         "critical_abs_z": 3, "horizon": "1w"},
        {"id": "lvl", "type": "level", "instrument": "OAT_BUND", "above": 80, "severity": "critical"},
        {"id": "chg", "type": "change", "instruments": ["class:rates"], "horizon": "mtd", "abs_above": 10},
    )
    assert rules[0].instruments == ("SX5E", "OAT_BUND") and rules[0].horizon.value == "1w"
    assert rules[1].severity is Severity.CRITICAL and rules[1].kind is RuleKind.LEVEL
    assert rules[2].condition == "abs_above" and rules[2].threshold == 10


@pytest.mark.parametrize("entry", [
    {"id": "a", "type": "magic", "instrument": "SX5E"},
    {"id": "a", "type": "zscore", "instrument": "SX5E", "min_abs_z": 2, "horizon": "mtd"},
    {"id": "a", "type": "zscore", "instrument": "SX5E", "min_abs_z": 3, "critical_abs_z": 2},
    {"id": "a", "type": "zscore", "instrument": "SX5E"},
    {"id": "a", "type": "level", "instrument": "SX5E", "abs_above": 1},
    {"id": "a", "type": "level", "instrument": "SX5E"},
    {"id": "a", "type": "level", "instrument": "SX5E", "above": 1, "below": 0},
    {"id": "a", "type": "change", "instruments": ["SX5E", "BUND_10Y"], "above": 1},
    {"id": "a", "type": "change", "instrument": "SX5E", "abs_above": -1},
    {"id": "a", "type": "level", "instrument": "GHOST", "above": 1},
    {"id": "a", "type": "stale", "instrument": "SX5E", "severity": "panic"},
    {"type": "stale", "instrument": "SX5E"},
])
def test_invalid_rules(entry):
    with pytest.raises(ReferentialError):
        _rules(entry)


def test_duplicate_ids_rejected():
    with pytest.raises(ReferentialError, match="duplicate"):
        _rules({"id": "a", "type": "stale", "instrument": "SX5E"},
               {"id": "a", "type": "stale", "instrument": "SX5E"})


def test_repository_rules_are_valid():
    from pathlib import Path

    from market_monitor.alerts import load_rules
    from market_monitor.referential import load_referential

    config = Path(__file__).resolve().parents[1] / "config"
    ref = load_referential(config / "instruments.yaml", config / "watchlists.yaml")
    assert len(load_rules(config / "alerts.yaml", ref)) >= 8


# ---------------------------------------------------------------------- engine
def test_zscore_rule_escalation_direction_and_stale():
    rules = _rules({"id": "z", "type": "zscore", "instruments": ["SX5E", "PROXY"],
                    "min_abs_z": 2, "critical_abs_z": 3.5, "direction": "down"})
    table = _table(SX5E=_row("Euro Stoxx 50", chg_1d=-3.2, z_1d=-4.0),
                   PROXY=_row("Proxy", chg_1d=2.5, z_1d=3.0))
    report = evaluate(rules, table, AS_OF)
    assert [(a.instrument_id, a.severity) for a in report.alerts] == [("SX5E", Severity.CRITICAL)]
    assert report.alerts[0].message == (
        "Euro Stoxx 50 : mouvement de 4,0 écarts-types sur la séance (−3,20\u202f%)")
    table.loc["SX5E", "stale"] = True
    assert not evaluate(rules, table, AS_OF).alerts


def test_level_rule_and_crossing():
    above = _rules({"id": "l", "type": "level", "instrument": "OAT_BUND", "above": 80})
    cross = _rules({"id": "c", "type": "level", "instrument": "OAT_BUND", "above": 80, "cross": True})
    row = _row("Spread OAT-Bund", quote="spread", unit="bp", asset_class="rates", level=82.0, ref_1d=78.0)
    table = _table(OAT_BUND=row)
    assert evaluate(above, table, AS_OF).alerts[0].message == (
        "Spread OAT-Bund : au-dessus de 80,0\u202fpb (82,0\u202fpb)")
    assert evaluate(cross, table, AS_OF).alerts[0].message == (
        "Spread OAT-Bund : franchit 80,0\u202fpb à la hausse (82,0\u202fpb)")
    table.loc["OAT_BUND", "ref_1d"] = 81.0  # already above before: no crossing
    assert not evaluate(cross, table, AS_OF).alerts
    below = _rules({"id": "b", "type": "level", "instrument": "OAT_BUND", "below": 0, "cross": True})
    table.loc["OAT_BUND", ["level", "ref_1d"]] = [-1.0, 2.0]
    assert "à la baisse" in evaluate(below, table, AS_OF).alerts[0].message


@pytest.mark.parametrize("condition,value,fires", [
    ({"above": 3}, 3.5, True), ({"above": 3}, 2.0, False),
    ({"below": -3}, -3.5, True), ({"abs_above": 3}, -3.5, True), ({"abs_above": 3}, 2.9, False),
])
def test_change_rule(condition, value, fires):
    rules = _rules({"id": "c", "type": "change", "instrument": "SX5E", "horizon": "1w", **condition})
    report = evaluate(rules, _table(SX5E=_row("Euro Stoxx 50", chg_1w=value)), AS_OF)
    assert bool(report.alerts) is fires


def test_change_message_and_stale_rule():
    rules = _rules({"id": "c", "type": "change", "instrument": "BUND_10Y", "abs_above": 10},
                   {"id": "s", "type": "stale", "instrument": "SX5E", "severity": "info"})
    table = _table(BUND_10Y=_row("Bund 10Y", quote="yield", unit="bp", asset_class="rates", chg_1d=-12.4),
                   SX5E=_row("Euro Stoxx 50", stale=True, level_date=pd.Timestamp("2026-09-25")))
    report = evaluate(rules, table, AS_OF)
    messages = [a.message for a in report.alerts]
    assert messages == ["Bund 10Y : −12,4\u202fpb sur 1J (seuil ±10,0\u202fpb)",
                        "Euro Stoxx 50 : dernière cotation le 25/09"]
    assert math.isnan(report.alerts[1].value)


def test_report_sorting_counts_and_missing_data():
    rules = _rules({"id": "i", "type": "stale", "instruments": ["SX5E", "PROXY"], "severity": "info"},
                   {"id": "w", "type": "level", "instrument": "SX5E", "above": 50, "severity": "critical"})
    table = _table(SX5E=_row("Euro Stoxx 50", stale=True), PROXY=_row("Proxy", level=math.nan))
    report = evaluate(rules, table, AS_OF)
    assert [a.rule_id for a in report.alerts] == ["w", "i"]  # most severe first
    assert report.count(Severity.CRITICAL) == 1 and report.max_severity is Severity.CRITICAL
    assert report.not_evaluated == {"i": ["PROXY"]}
    assert list(report.frame().columns)[:3] == ["severity", "rule", "instrument"]


# ----------------------------------------------------------------- integration
def _monitor():
    bbg, free = _providers()
    return MarketMonitor(MarketDataService([bbg, free]), REF)


RULES_RAW = {"rules": [
    {"id": "lvl", "name": "Spread OAT-Bund > 50 pb", "type": "level", "instrument": "OAT_BUND",
     "above": 50, "severity": "critical"},
    {"id": "z", "type": "zscore", "instruments": ["*"], "min_abs_z": 0.5},
]}


def test_monitor_alerts_end_to_end():
    report = _monitor().alerts(rules_from_dict(RULES_RAW, REF), AS_OF)
    assert report.alerts[0].rule_id == "lvl" and report.alerts[0].severity is Severity.CRITICAL
    assert "70,0\u202fpb" in report.alerts[0].message


def test_export_carries_alerts_to_text_and_excel(tmp_path):
    data = build_daily_macro(_monitor(), layout_from_dict(LAYOUT_RAW, REF), AS_OF,
                             rules_from_dict(RULES_RAW, REF))
    text = render_text(data)
    assert "Alertes du jour (usage interne)\n- [Critique] OAT-Bund : au-dessus de 50,0" in text
    from openpyxl import load_workbook
    write_excel(data, tmp_path / "x.xlsx")
    wb = load_workbook(tmp_path / "x.xlsx")
    assert wb.sheetnames == ["Daily macro", "Mouvements", "Texte", "Alertes", "Données"]
    assert wb["Alertes"]["A2"].value == "Critique"


def test_cli_alerts_exit_code_and_json(tmp_path, monkeypatch, capsys):
    (tmp_path / "alerts.yaml").write_text(yaml.safe_dump(RULES_RAW, allow_unicode=True), encoding="utf-8")
    config = tmp_path / "config.yaml"
    config.write_text("alerts:\n  rules: alerts.yaml\n", encoding="utf-8")
    monkeypatch.setattr(MarketMonitor, "from_settings", classmethod(lambda cls, s: _monitor()))
    out = tmp_path / "alerts.json"
    args = ["--config", str(config), "alerts", "--as-of", "2026-10-05", "--json", str(out)]
    assert main(args) == 0
    assert main([*args, "--fail-on", "critical"]) == 3
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload[0]["severity"] == "critical" and "[Critique]" in capsys.readouterr().out
