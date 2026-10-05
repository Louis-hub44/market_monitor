import math
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from market_monitor.cli import main
from market_monitor.data.service import MarketDataService
from market_monitor.export import (
    build_daily_macro,
    export_daily_macro,
    png_bytes,
    render_text,
    write_excel,
)
from market_monitor.export.excel import DATA_SHEET, change_formula, level_format
from market_monitor.export.layout import MoversRule, layout_from_dict
from market_monitor.export.text import move_phrase, section_line
from market_monitor.monitor import MarketMonitor
from market_monitor.referential import ReferentialError
from tests.test_history_monitor import REF, _providers

AS_OF = date(2026, 10, 5)
LAYOUT_RAW = {
    "title": "Revue test",
    "sections": [{"title": "Indices", "instruments": ["SX5E"]},
                 {"title": "Taux", "instruments": ["BUND_10Y", "OAT_BUND"]}],
    "columns": ["chg_1d", "chg_ytd"],
    "movers": {"universe": "all", "count": 3, "min_abs_z": 0.0, "strong_z": 2.5},
}


def _monitor():
    bbg, free = _providers()
    return MarketMonitor(MarketDataService([bbg, free]), REF)


def _data():
    return build_daily_macro(_monitor(), layout_from_dict(LAYOUT_RAW, REF), AS_OF)


# ---------------------------------------------------------------------- layout
def test_layout_resolves_watchlist_universe():
    layout = layout_from_dict(LAYOUT_RAW, REF)
    assert layout.section_ids == ("SX5E", "BUND_10Y", "OAT_BUND")
    assert layout.movers.universe == REF.ids  # watchlist "all" = "*"


@pytest.mark.parametrize("patch", [
    {"sections": []},
    {"columns": ["chg_2y"]},
    {"sections": [{"title": "X", "instruments": ["GHOST"]}]},
    {"file_prefix": "bad name"},
    {"movers": {"count": -1}},
])
def test_layout_validation(patch):
    with pytest.raises(ReferentialError):
        layout_from_dict({**LAYOUT_RAW, **patch}, REF)


# ------------------------------------------------------------------------ text
@pytest.mark.parametrize("args,expected", [
    ((-4.95, "pct", "price", -4.2, 2.5), "forte baisse de 4,95\u202f%"),
    ((1.2, "pct", "price", 1.0, 2.5), "hausse de 1,20\u202f%"),
    ((-5.0, "bp", "yield", -1.0, 2.5), "détente de 5,0\u202fpb"),
    ((4.0, "bp", "yield", math.nan, 2.5), "tension de 4,0\u202fpb"),
    ((-8.5, "bp", "spread", -5.4, 2.5), "fort resserrement de 8,5\u202fpb"),
    ((2.1, "abs", "vol", 0.5, 2.5), "hausse de 2,10\u202fpts"),
    ((0.001, "pct", "price", 0.0, 2.5), "stable"),
    ((math.nan, "pct", "price", 0.0, 2.5), "variation indisponible"),
])
def test_move_phrase(args, expected):
    assert move_phrase(*args) == expected


def test_section_line_marks_stale():
    row = pd.Series({"name": "Bitcoin", "level": 123627.0, "asset_class": "crypto", "quote": "price",
                     "change_unit": "pct", "chg_1d": -1.29, "chg_ytd": -11.26, "stale": True,
                     "level_date": pd.Timestamp("2026-09-25")})
    assert section_line(row, ("chg_1d", "chg_ytd")) == (
        "Bitcoin : 123\u202f627 (1J −1,29\u202f%, YTD −11,26\u202f%, cotation du 25/09)")


def test_render_text_structure():
    data = _data()
    text = render_text(data)
    assert text.startswith("Revue test, arrêté au lundi 5 octobre 2026")
    assert "Mouvements marquants" in text and "\nIndices\n- Euro Stoxx 50 : " in text
    assert "\nTaux\n- Bund 10Y : 2,700\u202f% (1J " in text
    assert "Sources : BBG, Libre." in text  # "derived" is not a source


# --------------------------------------------------------------------- builder
def test_builder_slices_one_computation():
    data = _data()
    assert [t for t, _ in data.sections] == ["Indices", "Taux"]
    assert list(data.sections[1][1].index) == ["BUND_10Y", "OAT_BUND"]
    assert len(data.movers) <= 3 and data.movers["z_1d"].abs().is_monotonic_decreasing
    assert data.universe.loc["BUND_10Y", "bp_factor"] == 100.0


def test_movers_threshold():
    from market_monitor.export.builder import select_movers
    table = pd.DataFrame({"z_1d": [0.5, -3.0, 2.0], "stale": [False] * 3}, index=list("ABC"))
    rule = MoversRule(universe=("A", "B", "C"), count=5, min_abs_z=1.5)
    assert list(select_movers(table, rule).index) == ["B", "C"]


# ----------------------------------------------------------------------- excel
def test_change_formulas_and_formats():
    assert change_formula("pct", 5, "1d") == (
        "=IF(OR(OR('Données'!$G$5=\"\",'Données'!$I$5=\"\"),N('Données'!$I$5)<=0),\"\","
        "'Données'!$G$5/'Données'!$I$5-1)")
    assert change_formula("bp", 7, "ytd").endswith("*'Données'!$F$7)")
    assert level_format("rates", "yield", 2.7) == '0.000" %"'
    assert level_format("fx", "price", 147.0) == "#,##0.00"


def test_excel_workbook_is_consistent(tmp_path):
    data = _data()
    path = tmp_path / "dm.xlsx"
    write_excel(data, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Daily macro", "Mouvements", "Texte", DATA_SHEET]
    macro, raw = wb["Daily macro"], wb[DATA_SHEET]
    assert macro["A1"].value == "Revue test"
    assert macro["A6"].value.startswith("='Données'!$B$")  # first instrument row
    row = int(macro["A6"].value.rsplit("$", 1)[1])
    level, ref_1d = raw[f"G{row}"].value, raw[f"I{row}"].value
    sx5e = data.universe.loc["SX5E"]
    assert level / ref_1d - 1 == pytest.approx(sx5e["chg_1d"] / 100)  # what the formula computes
    assert macro["C6"].number_format == "+0.00%;-0.00%;0.00%"
    assert all(c.font.name == "Arial" for r in macro.iter_rows() for c in r if c.value is not None)


# ------------------------------------------------------------------ png / files
def test_png_and_full_export(tmp_path):
    data = _data()
    assert png_bytes(data, dpi=50)[:8] == b"\x89PNG\r\n\x1a\n"
    result = export_daily_macro(_monitor(), layout_from_dict(LAYOUT_RAW, REF), tmp_path / "out", AS_OF)
    assert {p.name for p in (result.excel, result.png, result.text)} == {
        "daily_macro_2026-10-05.xlsx", "daily_macro_2026-10-05.png", "daily_macro_2026-10-05.txt"}
    assert result.text.read_text(encoding="utf-8").startswith("Revue test")


def test_cli_daily_macro(tmp_path, monkeypatch, capsys):
    import yaml

    (tmp_path / "dm.yaml").write_text(yaml.safe_dump(LAYOUT_RAW, allow_unicode=True), encoding="utf-8")
    config = tmp_path / "config.yaml"
    config.write_text("export:\n  layout: dm.yaml\n  output_dir: out\n", encoding="utf-8")
    monkeypatch.setattr(MarketMonitor, "from_settings", classmethod(lambda cls, s: _monitor()))
    assert main(["--config", str(config), "daily-macro", "--as-of", "2026-10-05"]) == 0
    assert (tmp_path / "out" / "daily_macro_2026-10-05.xlsx").is_file()
    assert "Revue test" in capsys.readouterr().out
