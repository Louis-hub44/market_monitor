"""Excel export (openpyxl).

* ``Daily macro`` : the bands, ready to paste (levels and changes are **formulas** on
  ``Données``, so every figure can be audited and recomputes if a level is corrected);
* ``Mouvements``  : the largest moves of the session;
* ``Texte``       : the text version, one line per cell;
* ``Données``     : levels, reference levels and dates per horizon, z-scores, sources,
  second-source check and the points to check.

A published line that needs checking (suspect move, contract roll) is shaded amber in
*Daily macro*, with the reason in a cell comment.

Percent changes are stored as fractions (``0.0182`` displayed ``+1.82%``), bp and
point changes as numbers with a unit in the number format.
"""

from __future__ import annotations

import math
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from market_monitor.alerts import AlertReport
from market_monitor.analytics.performance import Horizon
from market_monitor.analytics.quality import needs_check
from market_monitor.export.builder import DailyMacroData
from market_monitor.export.display import DEFAULT_STYLE, NumberStyle
from market_monitor.export.text import SEVERITY_LABELS, mover_line, render_text
from market_monitor.text_format import (
    ASSET_CLASS_LABELS,
    HORIZON_LABELS,
    NNBSP,
    SOURCE_LABELS,
    french_date,
)

FONT = "Arial"
NAVY, MUTED, ACCENT, UP, DOWN = "1B2835", "6B7B8C", "F0A830", "1E7F6F", "C4553F"
DATA_SHEET = "Données"
CHANGE_FORMATS = {
    "pct": '+0.00%;-0.00%;0.00%',
    "bp": '+0.0" pb";-0.0" pb";0.0" pb"',
    "abs": '+0.00" pts";-0.00" pts";0.00" pts"',
}
DATE_FORMAT = "dd/mm/yyyy"
DATA_COLUMNS = (
    ["Code", "Instrument", "Classe", "Cotation", "Convention", "Facteur pb", "Niveau", "Date"]
    + [label for h in Horizon
       for label in (f"Réf. {HORIZON_LABELS[h]}", f"Date réf. {HORIZON_LABELS[h]}")]
    + ["z 1J", "z 1S", "Source", "Proxy", "Périmé"]
    + ["Source effective", "Contrôle 2e source", "Suspect", "Roll 1J", "À vérifier"]
)
CHECK_FILL = "FCE9C8"  # light amber: line to check before publication
COL = {name: get_column_letter(k + 1) for k, name in enumerate(DATA_COLUMNS)}


def write_excel(data: DailyMacroData, target: str | Path | BinaryIO) -> None:
    """Write the workbook to a path or a binary buffer."""
    wb = Workbook()
    rows = _data_sheet(wb.create_sheet(DATA_SHEET), data.universe, data.checks)
    _macro_sheet(wb.active, data, rows)
    _movers_sheet(wb.create_sheet("Mouvements"), data, rows)
    _text_sheet(wb.create_sheet("Texte"), render_text(data))
    if data.alerts is not None:
        _alerts_sheet(wb.create_sheet("Alertes"), data.alerts)
    wb.move_sheet(DATA_SHEET, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(DATA_SHEET))
    wb.save(target)


def excel_bytes(data: DailyMacroData) -> bytes:
    buffer = BytesIO()
    write_excel(data, buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------- formulas
def ref(column: str, row: int) -> str:
    return f"'{DATA_SHEET}'!${COL[column]}${row}"


def change_formula(change_unit: str, row: int, horizon: str) -> str:
    """Change of ``Niveau`` vs ``Réf. <h>`` under the instrument convention (blank if missing)."""
    level, base = ref("Niveau", row), ref(f"Réf. {HORIZON_LABELS[horizon]}", row)
    guard = f'OR({level}="",{base}="")'
    if change_unit == "pct":
        return f'=IF(OR({guard},N({base})<=0),"",{level}/{base}-1)'
    if change_unit == "bp":
        return f'=IF({guard},"",({level}-{base})*{ref("Facteur pb", row)})'
    return f'=IF({guard},"",{level}-{base})'


def number_format(decimals: int, suffix: str = "", *, thousands: bool = True, sign: bool = False,
                  percent: bool = False) -> str:
    """Excel number format: ``(2, " $")`` -> ``#,##0.00" $"``; signed variants for changes."""
    core = ("#,##0" if thousands else "0") + (f".{'0' * decimals}" if decimals else "")
    unit = "%" if percent else (f'"{suffix}"' if suffix else "")
    if not sign:
        return f"{core}{unit}"
    return f"+{core}{unit};-{core}{unit};{core}{unit}"


def band_level_format(r: pd.Series, style: NumberStyle) -> str:
    """Level format of a band line (decimals / suffix from the layout)."""
    return number_format(int(r["level_decimals"]), str(r["level_suffix"]),
                         thousands=style.thousands_separator)


def band_change_format(r: pd.Series, style: NumberStyle) -> str:
    unit, decimals = str(r["display_unit"]), int(r["change_decimals"])
    if unit == "pct":
        return number_format(decimals, thousands=False, sign=True, percent=True)
    label = style.unit_label(unit)
    return number_format(decimals, label if style.compact_units else f" {label}",
                         thousands=False, sign=True)


def level_format(asset_class: str, quote: str, level: float) -> str:
    if quote == "yield":
        return '0.000" %"'
    if quote == "spread":
        return '0.0" pb"'
    if asset_class == "fx":
        return "0.0000" if abs(level) < 20 else "#,##0.00"
    return "#,##0" if asset_class == "crypto" else "#,##0.00"


# ------------------------------------------------------------------- sheets
def _data_sheet(ws: Worksheet, table: pd.DataFrame,
                checks: dict[str, list[str]] | None = None) -> dict[str, int]:
    """Write the audit table; returns ``id -> row``."""
    ws.append(DATA_COLUMNS)
    rows: dict[str, int] = {}
    for k, (instrument_id, r) in enumerate(table.iterrows(), start=2):
        values: list[Any] = [
            instrument_id, r["name"], ASSET_CLASS_LABELS.get(r["asset_class"], r["asset_class"]),
            r["quote"], r["change_unit"], r["bp_factor"], r["level"], r["level_date"],
        ]
        for h in Horizon:
            values += [r[f"ref_{h}"], r[f"ref_date_{h}"]]
        values += [r["z_1d"], r["z_1w"], SOURCE_LABELS.get(r["source"], r["source"]),
                   r["proxy"], "oui" if r["stale"] else "non"]
        values += [r.get("origin"), r.get("cross_check") or "non fait",
                   "oui" if r.get("suspect") else "non", "oui" if r.get("roll_1d") else "non",
                   _xl_text(" ; ".join((checks or {}).get(str(instrument_id), [])))]
        ws.append([_cell(v) for v in values])
        rows[str(instrument_id)] = k
        for column in [c for c in DATA_COLUMNS if c.startswith("Date")]:
            ws[f"{COL[column]}{k}"].number_format = DATE_FORMAT
    _style_header(ws, 1, len(DATA_COLUMNS))
    _finish(ws, widths={"A": 14, "B": 26, COL["Source effective"]: 22, COL["À vérifier"]: 60})
    ws.freeze_panes = "C2"
    return rows


def _macro_sheet(ws: Worksheet, data: DailyMacroData, rows: dict[str, int]) -> None:
    ws.title = "Daily macro"
    ws["A1"] = data.title
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    ws["A2"] = f"Arrêté au {french_date(data.as_of)}"
    ws["A2"].font = Font(name=FONT, italic=True, color=MUTED)
    horizons = [HORIZON_LABELS[c.removeprefix("chg_")] for c in data.columns]
    headers = ["Instrument", "Niveau", *horizons, "Date"]
    line = 4
    for title, table in data.sections:
        ws.cell(line, 1, title).font = Font(name=FONT, bold=True, size=11, color=NAVY)
        for c in range(1, len(headers) + 1):
            ws.cell(line, c).border = Border(bottom=Side(style="medium", color=ACCENT))
        line += 1
        for c, header in enumerate(headers, start=1):
            ws.cell(line, c, header)
        _style_header(ws, line, len(headers))
        first = line + 1
        for instrument_id, r in table.iterrows():
            line += 1
            _macro_row(ws, line, rows[str(instrument_id)], r, data.columns,
                       data.checks.get(str(instrument_id), []), data.style)
        _sign_colours(ws, f"C{first}:{get_column_letter(2 + len(data.columns))}{line}")
        line += 2
    _finish(ws, widths={"A": 28, "B": 13} | {get_column_letter(3 + k): 11 for k in range(len(headers))})


def _macro_row(ws: Worksheet, line: int, data_row: int, r: pd.Series, columns: tuple[str, ...],
               checks: list[str] | None = None, style: NumberStyle = DEFAULT_STYLE) -> None:
    banded = "level_text" in r
    label = r["label"] if banded and r["label"] != r["name"] else None
    ws.cell(line, 1, _xl_text(label) if label else f"={ref('Instrument', data_row)}")
    level = ws.cell(line, 2, f'=IF({ref("Niveau", data_row)}="","",{ref("Niveau", data_row)})')
    level.number_format = (band_level_format(r, style) if banded
                           else level_format(r["asset_class"], r["quote"], _num(r["level"])))
    unit = r["display_unit"] if banded else r["change_unit"]
    for k, column in enumerate(columns, start=3):
        cell = ws.cell(line, k, change_formula(unit, data_row, column.removeprefix("chg_")))
        cell.number_format = band_change_format(r, style) if banded else CHANGE_FORMATS[unit]
    date_cell = ws.cell(line, 3 + len(columns), f'=IF({ref("Date", data_row)}="","",{ref("Date", data_row)})')
    date_cell.number_format = "dd/mm"
    if r["stale"]:
        for c in range(1, 4 + len(columns)):
            ws.cell(line, c).font = Font(name=FONT, italic=True, color=MUTED)
    if needs_check(r):
        for c in range(1, 4 + len(columns)):
            ws.cell(line, c).fill = PatternFill("solid", fgColor=CHECK_FILL)
        comment = Comment(_xl_text("À vérifier avant diffusion :\n- " + "\n- ".join(checks or [])),
                          "Market Monitor")
        comment.width, comment.height = 320, 110
        ws.cell(line, 1).comment = comment


def _movers_sheet(ws: Worksheet, data: DailyMacroData, rows: dict[str, int]) -> None:
    headers = ["Rang", "Instrument", "Variation 1J", "z 1J", "Commentaire"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    rule = data.movers_rule
    for rank, (instrument_id, r) in enumerate(data.movers.iterrows(), start=1):
        line, data_row = rank + 1, rows[str(instrument_id)]
        ws.cell(line, 1, rank)
        ws.cell(line, 2, f"={ref('Instrument', data_row)}")
        change = ws.cell(line, 3, change_formula(r["change_unit"], data_row, "1d"))
        change.number_format = CHANGE_FORMATS[r["change_unit"]]
        ws.cell(line, 4, f"={ref('z 1J', data_row)}").number_format = "+0.0;-0.0;0.0"
        ws.cell(line, 5, _xl_text(mover_line(r, rule.strong_z, show_zscore=False)))
    if not data.movers.empty:
        _sign_colours(ws, f"C2:D{len(data.movers) + 1}")
    _finish(ws, widths={"A": 7, "B": 28, "C": 13, "D": 8, "E": 70})


SEVERITY_FILLS = {"critical": "F6D5CF", "warning": "FCE9C8", "info": "DCEBF7"}


def _alerts_sheet(ws: Worksheet, report: AlertReport) -> None:
    headers = ["Gravité", "Règle", "Instrument", "Message", "Valeur", "Seuil", "Date"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for line, a in enumerate(report.alerts, start=2):
        values = [SEVERITY_LABELS[a.severity.value], a.rule, a.instrument, _xl_text(a.message),
                  _cell(a.value), _cell(a.threshold), _cell(a.level_date)]
        for c, value in enumerate(values, start=1):
            ws.cell(line, c, value)
        ws.cell(line, 1).fill = PatternFill("solid", fgColor=SEVERITY_FILLS[a.severity.value])
        ws.cell(line, 7).number_format = "dd/mm/yyyy"
    if not report.alerts:
        ws.cell(2, 1, "Aucune alerte déclenchée.")
    _finish(ws, widths={"A": 11, "B": 34, "C": 24, "D": 70, "E": 10, "F": 10, "G": 12})


def _text_sheet(ws: Worksheet, text: str) -> None:
    for k, line in enumerate(text.splitlines(), start=1):
        ws.cell(k, 1, _xl_text(line))
    _finish(ws, widths={"A": 110})


# ------------------------------------------------------------------ helpers
def _cell(value: Any) -> Any:
    """Excel-safe value: NaN / NaT -> blank, Timestamp -> datetime."""
    if value is None or (isinstance(value, float) and math.isnan(value)) or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _xl_text(text: str) -> str:
    """Excel fonts do not all have U+202F: use a regular no-break space."""
    return text.replace(NNBSP, "\u00a0")


def _num(value: Any) -> float:
    return 0.0 if _cell(value) is None else float(value)


def _style_header(ws: Worksheet, row: int, n_columns: int) -> None:
    for c in range(1, n_columns + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name=FONT, bold=True, color=MUTED, size=9)
        cell.alignment = Alignment(horizontal="left" if c == 1 else "right")
        cell.border = Border(bottom=Side(style="thin", color=MUTED))


def _sign_colours(ws: Worksheet, cell_range: str) -> None:
    """Dynamic font colour by sign (follows the formulas if a level is corrected)."""
    ws.conditional_formatting.add(
        cell_range, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, color=UP)))
    ws.conditional_formatting.add(
        cell_range, CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color=DOWN)))


def _finish(ws: Worksheet, widths: dict[str, float]) -> None:
    """Arial everywhere (keeping weights / colours already set) and column widths."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold, italic=cell.font.italic,
                                 color=cell.font.color, size=cell.font.size or 10)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.showGridLines = False
